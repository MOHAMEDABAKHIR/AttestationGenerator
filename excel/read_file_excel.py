from __future__ import annotations

import os
from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook


@dataclass
class ExcelPreview:
    headers: list[str]
    rows: list[list[Any]]
    header_row: int
    company_column: int | None


class ExcelReader:
    """
    Service responsable uniquement de la lecture et de l'analyse
    des fichiers Excel.

    L'interface graphique ne doit pas manipuler directement openpyxl.
    """

    def __init__(self, file_path: str):
        if not file_path:
            raise ValueError("Aucun fichier Excel n'a été fourni.")

        if not os.path.exists(file_path):
            raise FileNotFoundError(
                f"Le fichier Excel n'existe pas : {file_path}"
            )

        self.file_path = file_path

        self.workbook = load_workbook(
            file_path,
            data_only=True,
            read_only=True
        )

    # ------------------------------------------------------------------
    # INFORMATIONS FICHIER
    # ------------------------------------------------------------------

    def get_sheet_names(self) -> list[str]:
        return self.workbook.sheetnames

    def get_sheet(self, sheet_name: str):
        if sheet_name not in self.workbook.sheetnames:
            raise ValueError(
                f"La feuille '{sheet_name}' n'existe pas."
            )

        return self.workbook[sheet_name]

    # ------------------------------------------------------------------
    # LECTURE DES LIGNES
    # ------------------------------------------------------------------

    def get_row_values(
        self,
        sheet_name: str,
        row_number: int
    ) -> list[Any]:

        sheet = self.get_sheet(sheet_name)

        if row_number < 1 or row_number > sheet.max_row:
            return []

        values = []

        for cell in sheet[row_number]:
            values.append(cell.value)

        return values

    def get_non_empty_rows(
        self,
        sheet_name: str,
        max_preview_rows: int = 200
    ) -> list[tuple[int, list[Any]]]:

        sheet = self.get_sheet(sheet_name)

        result = []

        max_row = min(sheet.max_row, max_preview_rows)

        for row_number in range(1, max_row + 1):

            values = [
                cell.value
                for cell in sheet[row_number]
            ]

            if any(
                value is not None and str(value).strip() != ""
                for value in values
            ):
                result.append(
                    (row_number, values)
                )

        return result

    # ------------------------------------------------------------------
    # DETECTION DES HEADERS
    # ------------------------------------------------------------------

    def get_header_candidates(
        self,
        sheet_name: str,
        max_rows: int = 100
    ) -> list[tuple[int, list[Any]]]:

        """
        Retourne les lignes qui contiennent au moins
        2 cellules non vides.

        On ne décide pas automatiquement du header :
        l'utilisateur le choisit manuellement.
        """

        sheet = self.get_sheet(sheet_name)

        candidates = []

        max_row = min(sheet.max_row, max_rows)

        for row_number in range(1, max_row + 1):

            values = [
                cell.value
                for cell in sheet[row_number]
            ]

            non_empty_count = sum(
                1
                for value in values
                if value is not None
                and str(value).strip() != ""
            )

            if non_empty_count >= 2:
                candidates.append(
                    (row_number, values)
                )

        return candidates

    # ------------------------------------------------------------------
    # HEADERS
    # ------------------------------------------------------------------

    def get_headers(
        self,
        sheet_name: str,
        header_row: int
    ) -> list[str]:

        values = self.get_row_values(
            sheet_name,
            header_row
        )

        headers = []

        for index, value in enumerate(values):

            if value is None or str(value).strip() == "":
                headers.append(
                    f"Colonne {index + 1}"
                )
            else:
                headers.append(
                    str(value).strip()
                )

        return headers

    # ------------------------------------------------------------------
    # ENTREPRISES
    # ------------------------------------------------------------------

    def get_column_values(
        self,
        sheet_name: str,
        start_row: int,
        column_index: int
    ) -> list[str]:

        sheet = self.get_sheet(sheet_name)

        values = []

        for row_number in range(
            start_row + 1,
            sheet.max_row + 1
        ):

            value = sheet.cell(
                row=row_number,
                column=column_index
            ).value

            if value is None:
                continue

            value = str(value).strip()

            if not value:
                continue

            values.append(value)

        return values

    def get_unique_company_names(
        self,
        sheet_name: str,
        header_row: int,
        company_column: int
    ) -> list[str]:

        values = self.get_column_values(
            sheet_name,
            header_row,
            company_column
        )

        # Suppression des doublons en conservant l'ordre
        unique_values = list(
            dict.fromkeys(values)
        )

        return unique_values

    # ------------------------------------------------------------------
    # PREVIEW
    # ------------------------------------------------------------------

    def get_preview(
        self,
        sheet_name: str,
        header_row: int,
        max_rows: int = 15
    ) -> ExcelPreview:

        sheet = self.get_sheet(sheet_name)

        headers = self.get_headers(
            sheet_name,
            header_row
        )

        rows = []

        start = header_row + 1
        end = min(
            sheet.max_row,
            start + max_rows - 1
        )

        for row_number in range(start, end + 1):

            values = [
                sheet.cell(
                    row=row_number,
                    column=column
                ).value
                for column in range(
                    1,
                    sheet.max_column + 1
                )
            ]

            if not any(
                value is not None
                and str(value).strip() != ""
                for value in values
            ):
                continue

            rows.append(values)

        return ExcelPreview(
            headers=headers,
            rows=rows,
            header_row=header_row,
            company_column=None
        )


# ----------------------------------------------------------------------
# COMPATIBILITÉ AVEC TON ANCIEN CODE
# ----------------------------------------------------------------------

def column_letter_to_number(letter: str) -> int:
    number = 0

    for char in letter.upper():
        number = (
            number * 26
            + ord(char)
            - ord("A")
            + 1
        )

    return number


def extractFileSheets(file) -> str:
    return ",".join(
        sheet.title
        for sheet in file
    )


def extractSheetSize(file, sheet_name: str) -> str:
    sheet = file[sheet_name]

    return f"{sheet.max_row}*{sheet.max_column}"


def findTableStructure(
    file,
    sheet_name,
    table_header_content
):
    """
    Ancienne fonction conservée pour éviter de casser
    d'éventuels appels existants.

    La nouvelle interface ne dépend plus de cette fonction.
    """

    sheet = file[sheet_name]

    expected = {
        str(value).strip().lower()
        for value in table_header_content
    }

    best_row = None
    best_column = None
    best_score = 0

    for row_number in range(
        1,
        sheet.max_row + 1
    ):

        for column_number in range(
            1,
            sheet.max_column + 1
        ):

            value = sheet.cell(
                row=row_number,
                column=column_number
            ).value

            if value is None:
                continue

            normalized = str(value).strip().lower()

            if normalized not in expected:
                continue

            score = 0

            for header_index in range(
                column_number,
                min(
                    sheet.max_column + 1,
                    column_number + len(table_header_content)
                )
            ):

                current = sheet.cell(
                    row=row_number,
                    column=header_index
                ).value

                if current is None:
                    continue

                if str(current).strip().lower() in expected:
                    score += 1

            if score > best_score:
                best_score = score
                best_row = row_number
                best_column = column_number

    if best_row is None:
        return None

    from openpyxl.utils import get_column_letter

    return [
        best_row,
        get_column_letter(best_column)
    ]


def writeTable(
    file_path,
    sheet_name,
    table_header_content
):

    reader = ExcelReader(file_path)

    structure = findTableStructure(
        reader.workbook,
        sheet_name,
        table_header_content
    )

    if structure is None:
        raise ValueError(
            "Impossible de détecter automatiquement "
            "la structure du tableau."
        )

    header_row = structure[0]
    start_column = column_letter_to_number(
        structure[1]
    )

    headers = reader.get_headers(
        sheet_name,
        header_row
    )

    end_column = min(
        len(headers),
        start_column - 1
        + len(table_header_content)
    )

    data = [headers]

    sheet = reader.get_sheet(sheet_name)

    for row_number in range(
        header_row + 1,
        sheet.max_row + 1
    ):

        values = []

        for column_number in range(
            start_column,
            end_column + 1
        ):

            values.append(
                sheet.cell(
                    row=row_number,
                    column=column_number
                ).value
            )

        if not any(
            value is not None
            and str(value).strip() != ""
            for value in values
        ):
            continue

        data.append(values)

    return data