from __future__ import annotations

import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from datetime import datetime

from docx import Document


class TemplateGenerator:
    """
    Générateur des attestations à partir de modèles Word.

    Le principe est volontairement simple :

        template Word
              ↓
        remplacement des variables
              ↓
        document Word final
    """

    def __init__(self, templates_dir: str):

        self.templates_dir = Path(
            templates_dir
        )

    # =========================================================
    # TEMPLATE
    # =========================================================

    def get_template_path(
        self,
        signer_quality: str,
        certificate_type: str
    ) -> Path:

        quality = self.normalize_quality(
            signer_quality
        )

        certificate = self.normalize_certificate_type(
            certificate_type
        )

        templates = {

            (
                "commissaire_aux_comptes",
                "avec_retard"
            ): "CAC_avec_retard.docx",

            (
                "commissaire_aux_comptes",
                "sans_retard"
            ): "CAC_sans_retard.docx",

            (
                "expert_comptable",
                "avec_retard"
            ): "EC_avec_retard.docx",

            (
                "expert_comptable",
                "sans_retard"
            ): "EC_sans_retard.docx",
        }

        filename = templates.get(
            (quality, certificate)
        )

        if not filename:
            raise ValueError(
                "Combinaison de template inconnue : "
                f"{signer_quality} / {certificate_type}"
            )

        path = self.templates_dir / filename

        if not path.exists():

            raise FileNotFoundError(
                f"Le modèle Word est introuvable :\n{path}"
            )

        return path

    # =========================================================
    # NORMALISATION
    # =========================================================

    @staticmethod
    def normalize_quality(
        quality: str
    ) -> str:

        value = (
            str(quality or "")
            .strip()
            .lower()
        )

        value = (
            value
            .replace("-", " ")
            .replace("_", " ")
        )

        if (
            "commissaire" in value
            or value == "cac"
        ):
            return "commissaire_aux_comptes"

        if (
            "expert" in value
            or value == "ec"
        ):
            return "expert_comptable"

        raise ValueError(
            f"Qualité de signataire inconnue : {quality}"
        )

    @staticmethod
    def normalize_certificate_type(certificate_type):
        value = str(certificate_type or "").strip().lower()

        # IMPORTANT :
        # tester "sans retard" AVANT "avec retard"

        if "sans retard" in value:
            return "sans_retard"

        if "avec retard" in value:
            return "avec_retard"

        raise ValueError(
            f"Type d'attestation inconnu : {certificate_type}"
        )

    # =========================================================
    # DOCX
    # =========================================================

    def replace_text_in_paragraph(
        self,
        paragraph,
        replacements: dict[str, str]
    ):

        """
        Remplace les variables dans un paragraphe.

        On travaille sur le texte global du paragraphe.
        """

        original = paragraph.text

        if not original:
            return

        updated = original

        for old, new in replacements.items():

            updated = updated.replace(
                old,
                str(new)
            )

        if updated == original:
            return

        # Remplacement simple.
        # Pour les templates finaux, on pourra ensuite
        # renforcer la conservation de certains styles.
        paragraph.text = updated

    def replace_text_in_table(
        self,
        table,
        replacements
    ):

        for row in table.rows:

            for cell in row.cells:

                for paragraph in cell.paragraphs:

                    self.replace_text_in_paragraph(
                        paragraph,
                        replacements
                    )

                for nested_table in cell.tables:

                    self.replace_text_in_table(
                        nested_table,
                        replacements
                    )

    def replace_everywhere(
        self,
        document,
        replacements
    ):

        # Paragraphes normaux

        for paragraph in document.paragraphs:

            self.replace_text_in_paragraph(
                paragraph,
                replacements
            )

        # Tableaux

        for table in document.tables:

            self.replace_text_in_table(
                table,
                replacements
            )

        # Headers / Footers

        for section in document.sections:

            for paragraph in section.header.paragraphs:

                self.replace_text_in_paragraph(
                    paragraph,
                    replacements
                )

            for paragraph in section.footer.paragraphs:

                self.replace_text_in_paragraph(
                    paragraph,
                    replacements
                )

    # =========================================================
    # GENERATION
    # =========================================================

    def generate(
        self,
        data: dict,
        output_directory: str
    ) -> str:

        signer_quality = data.get(
            "signer_quality",
            ""
        )

        certificate_type = data.get(
            "certificate_type",
            ""
        )

        template_path = self.get_template_path(
            signer_quality,
            certificate_type
        )

        output_dir = Path(
            output_directory
        )

        output_dir.mkdir(
            parents=True,
            exist_ok=True
        )

        company_name = data.get(
            "company",
            "Attestation"
        )

        safe_company_name = self.sanitize_filename(
            company_name
        )

        output_path = (
            output_dir
            / f"{safe_company_name}.docx"
        )

        # -----------------------------------------------------
        # Variables
        # -----------------------------------------------------

        replacements = self.build_replacements(
            data
        )

        # -----------------------------------------------------
        # Document
        # -----------------------------------------------------

        document = Document(
            str(template_path)
        )

        self.replace_everywhere(
            document,
            replacements
        )

        document.save(
            str(output_path)
        )

        return str(output_path)

    # =========================================================
    # REPLACEMENTS
    # =========================================================

    def build_replacements(
        self,
        data: dict
    ) -> dict[str, str]:

        company = data.get(
            "company",
            ""
        )

        period = data.get(
            "period",
            ""
        )

        gender = data.get(
            "gender",
            "Monsieur"
        )

        ceo = data.get(
            "ceo",
            ""
        )

        address = data.get(
            "address",
            ""
        )

        amount = data.get(
            "amount",
            ""
        )

        signature_date = data.get(
            "signature_date",
            ""
        )

        # -----------------------------------------------------
        # Gérant / Gérante
        # -----------------------------------------------------

        gerant_label = (
            "Gérante"
            if gender == "Madame"
            else "Gérant"
        )

        return {

            "COMPANY":
                company,

            "Company":
                company,

            "company":
                company,

            "PERIODE":
                period,

            "PERIODE DU 1er AVRIL AU 30 JUIN 2026":
                period,

            "Nom & prénom":
                ceo,

            "Monsieur Nom & prénom":
                f"{gender} {ceo}",

            "Gérant COMPANY":
                f"{gerant_label} {company}",

            "Gérant":
                gerant_label,

            "ZENITH MILLENIUM IMM 2 ETG 2 LOT ATTAWFIK S.":
                address,

            "123.026,05":
                str(amount),

            "30/07/2026":
                signature_date,

        }

    # =========================================================
    # FILENAME
    # =========================================================

    @staticmethod
    def sanitize_filename(
        filename: str
    ) -> str:

        filename = str(
            filename or "Attestation"
        ).strip()

        filename = re.sub(
            r'[<>:"/\\|?*]',
            "_",
            filename
        )

        filename = re.sub(
            r"\s+",
            " ",
            filename
        )

        return filename[:150]